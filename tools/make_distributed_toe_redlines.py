from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


OUT_DIR = Path("review_artifacts")
AUTHOR = "Codex draft"
STAMP = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def set_track_revisions(doc):
    settings = doc.settings.element
    track = settings.find(qn("w:trackRevisions"))
    if track is None:
        settings.append(OxmlElement("w:trackRevisions"))


def set_run_style(run):
    run.font.name = "Aptos"
    run.font.size = Pt(10)


def add_deleted(paragraph, text, change_id):
    deleted = OxmlElement("w:del")
    deleted.set(qn("w:id"), str(change_id))
    deleted.set(qn("w:author"), AUTHOR)
    deleted.set(qn("w:date"), STAMP)

    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "C00000")
    strike = OxmlElement("w:strike")
    props.append(color)
    props.append(strike)
    run.append(props)
    del_text = OxmlElement("w:delText")
    del_text.set(qn("xml:space"), "preserve")
    del_text.text = text
    run.append(del_text)
    deleted.append(run)
    paragraph._p.append(deleted)


def add_inserted(paragraph, text, change_id):
    inserted = OxmlElement("w:ins")
    inserted.set(qn("w:id"), str(change_id))
    inserted.set(qn("w:author"), AUTHOR)
    inserted.set(qn("w:date"), STAMP)

    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0070C0")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(color)
    props.append(underline)
    run.append(props)
    ins_text = OxmlElement("w:t")
    ins_text.set(qn("xml:space"), "preserve")
    ins_text.text = text
    run.append(ins_text)
    inserted.append(run)
    paragraph._p.append(inserted)


def add_review_paragraph(doc, parts, style=None):
    p = doc.add_paragraph(style=style)
    for mode, text in parts:
        if mode == "normal":
            run = p.add_run(text)
            set_run_style(run)
        elif mode == "delete":
            add_deleted(p, text, add_review_paragraph.change_id)
            add_review_paragraph.change_id += 1
        elif mode == "insert":
            add_inserted(p, text, add_review_paragraph.change_id)
            add_review_paragraph.change_id += 1
        else:
            raise ValueError(mode)
    return p


add_review_paragraph.change_id = 1


def add_source_note(doc, source_path):
    p = doc.add_paragraph()
    run = p.add_run(f"Source file: {source_path}")
    run.italic = True
    run.font.size = Pt(9)
    run.font.highlight_color = WD_COLOR_INDEX.GRAY_25


def build_doc(filename, title, source_path, changes):
    add_review_paragraph.change_id = 1
    doc = Document()
    set_track_revisions(doc)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    doc.add_heading(title, level=1)
    add_source_note(doc, source_path)
    doc.add_paragraph(
        "This review copy shows the proposed distributed TOE and microservices changes as tracked revisions. "
        "Blue underlined text is inserted text; red strikethrough text is deleted text."
    )
    for heading, paragraphs in changes:
        doc.add_heading(heading, level=2)
        for parts in paragraphs:
            add_review_paragraph(doc, parts)
    OUT_DIR.mkdir(exist_ok=True)
    doc.save(OUT_DIR / filename)


def source_block_as_insertions(source_path, start_marker, end_marker):
    text = Path(source_path).read_text(encoding="utf-8")
    start = text.index(start_marker)
    end = text.index(end_marker, start)
    block = text[start:end].strip()
    paragraphs = []
    for line in block.splitlines():
        if line.strip():
            paragraphs.append([("insert", line)])
    return paragraphs


base_changes = [
    (
        "New Section After Compliant TOE / Use Cases",
        [
            [
                ("insert", "Distributed and Microservices TOE Architectures"),
            ],
            [
                (
                    "insert",
                    "A TOE may consist of multiple separately deployed application components that collectively provide the TOE security functionality. Examples include server-agent products, clustered application deployments, and microservices-based applications composed of multiple application payloads. If a TOE is distributed across multiple TOE components, the ST shall identify each TOE component, describe the role of each TOE component, identify which components implement each claimed SFR, describe all communications between TOE components, and distinguish TOE components from operational environment components.",
                )
            ],
            [
                (
                    "insert",
                    "For containerized or microservices TOEs, the TOE consists of the application payloads and TOE-provided application components identified in the ST. The container orchestration platform, container runtime, operating system, service mesh infrastructure, ingress infrastructure, cluster networking, and platform-provided secret or configuration stores are part of the operational environment unless explicitly included in the TOE boundary. The PP does not require these operational environment components to be included in the TOE boundary solely because the TOE depends on them for execution, scheduling, networking, isolation, credential storage, configuration storage, or time services.",
                )
            ],
            [
                (
                    "insert",
                    "When the TOE relies on operational environment components to provide services used by the TOE, the ST shall identify the dependency and the guidance shall describe the required environmental configuration. Inter-component communication between TOE components shall be identified in the ST. Where the TOE claims conformance to a PP-Configuration that includes Server and Agent application modules, the ST shall use the module requirements to address authorization or registration of TOE components before communication is permitted and FTP_DIT_EXT.1 to address protection of data transmitted between TOE components. The ST shall provide an SFR allocation rationale that identifies whether each claimed requirement is satisfied by all TOE components, by applicable TOE components that perform the relevant function, by at least one TOE component, as a feature dependent requirement, or by an allowed operational environment dependency.",
                )
            ],
        ],
    )
]

server_changes = [
    (
        "TOE Overview",
        [
            [
                (
                    "normal",
                    "This is a Collaborative Protection Profile (cPP) Module whose Target of Evaluation (TOE) is Enterprise Server Applications. This PP-Module is compatible with the cPP for Application Software. ",
                )
            ],
            [
                (
                    "insert",
                    "For a distributed TOE, the Server Application is the TOE component, or set of TOE components, that provides management, coordination, policy, API-facing, or other server-side functionality for the TOE. In a microservices architecture, a Server Application component may be a service or application payload that coordinates, exposes, or controls TOE functionality. The container orchestration platform, container runtime, operating system, service mesh infrastructure, ingress infrastructure, cluster networking, and platform-provided secret or configuration stores are part of the operational environment unless explicitly included in the TOE boundary.",
                )
            ],
        ],
    ),
    (
        "New Distributed/Microservices Section",
        [
            [("insert", "Distributed and Microservices TOE Configurations")],
            [
                (
                    "insert",
                    "This PP-Module may be used in a PP-Configuration with the PP-Module for Agent Applications to evaluate distributed application software. Distributed application software includes server-agent deployments, clustered server deployments, and microservices architectures composed of multiple application payload components.",
                )
            ],
            [
                (
                    "insert",
                    "For a distributed TOE, the ST shall identify each TOE component, describe the role of each TOE component, identify which components implement each claimed SFR, and describe all communications between TOE components. The ST shall also distinguish TOE components from operational environment components. Operational environment components may include container orchestration, container runtimes, operating systems, service mesh infrastructure, ingress infrastructure, cluster networking, platform-provided secret or configuration stores, and other infrastructure services not explicitly included in the TOE boundary.",
                )
            ],
            [
                (
                    "insert",
                    "If the TOE relies on operational environment components for execution, scheduling, networking, isolation, credential storage, configuration storage, time services, or protection of inter-component communications, the ST shall identify the dependency and the guidance shall describe the required environmental configuration.",
                )
            ],
        ],
    ),
    (
        "SFR Applicability Updates",
        [
            [
                ("normal", "These SFRs apply "),
                ("delete", "if and only if an Agent Module is included in the evaluation."),
                (
                    "insert",
                    "when the TOE includes separately deployed TOE components that communicate with one another as part of a PP-Configuration that includes the Agent Module. For microservices architectures, these SFRs apply to the communication relationships between Server Application components and Agent Application components as those components are identified in the ST. The ST author should iterate these SFRs as needed for different component pairs or communication mechanisms.",
                ),
            ],
            [
                ("normal", "configuration of communication with "),
                ("delete", "Agent"),
                ("insert", "other TOE components"),
                ("normal", " according to FMT_SMF.1/Server and FTP_DIT_EXT.1"),
            ],
        ],
    ),
]

agent_changes = [
    (
        "TOE Overview",
        [
            [
                (
                    "normal",
                    "This is a Collaborative Protection Profile (cPP) Module whose Target of Evaluation (TOE) is Agent Applications. This PP-Module is compatible with the cPP for Application Software and collaborative PP-Module for Server Applications. ",
                )
            ],
            [
                (
                    "insert",
                    "For purposes of a PP-Configuration, an Agent Application is any separately deployed TOE component that communicates with another TOE component under the control, coordination, policy, enrollment, or trust relationship established by the TOE. This may include endpoint agents, worker services, peer services, microservice payloads, subordinate application services, or other application components that are identified as TOE components in the ST.",
                )
            ],
            [
                (
                    "insert",
                    "For containerized or microservices TOEs, the TOE consists of the application payload components identified in the ST. The container orchestration platform, container runtime, operating system, service mesh infrastructure, ingress infrastructure, cluster networking, and platform-provided secret or configuration stores are part of the operational environment unless explicitly included in the TOE boundary.",
                )
            ],
        ],
    ),
    (
        "New Distributed/Microservices Section",
        [
            [("insert", "Distributed and Microservices TOE Configurations")],
            [
                (
                    "insert",
                    "This PP-Module may be used in a PP-Configuration with the PP-Module for Server Applications to evaluate distributed application software. Distributed application software includes server-agent deployments, clustered server deployments, and microservices architectures composed of multiple application payload components.",
                )
            ],
            [
                (
                    "insert",
                    "The ST shall identify each Agent Application component, describe the role of each component, identify which claimed SFRs are implemented by each component, and describe all communications between Agent Application components and other TOE components. The ST shall also distinguish TOE components from operational environment components. If the TOE relies on operational environment components for execution, scheduling, networking, isolation, credential storage, configuration storage, time services, or protection of inter-component communications, the ST shall identify the dependency and the guidance shall describe the required environmental configuration.",
                )
            ],
        ],
    ),
    (
        "FCO Application Note",
        [
            [
                ("delete", "An Agent can communicate with a Server or another Agent."),
                (
                    "insert",
                    "An Agent can communicate with a Server, another Agent, or another separately deployed TOE component identified in the ST. In a microservices architecture, this may include communication between application payload services.",
                ),
                ("normal", " This SFR can be iterated if the registration method varies depending on which TOE components are communicating."),
            ]
        ],
    ),
]

config_changes = [
    (
        "Title and Overview",
        [
            [
                ("delete", "PP-Configuration for Enterprise Server Applications and Client Agent(s)"),
                ("insert", "PP-Configuration for Enterprise Server Applications and Agent/Application Component(s)"),
            ],
            [
                ("normal", "This PP-Configuration is for enterprise server applications and their "),
                ("delete", "client agent(s)."),
                (
                    "insert",
                    "agent or application component(s). It provides the enforceable PP-Configuration path for distributed application software, including server-agent deployments, clustered server deployments, and microservices architectures composed of multiple application payload components.",
                ),
            ],
        ],
    ),
    (
        "New Distributed/Microservices Section",
        [
            [("insert", "Distributed and Microservices TOE Architectures")],
            [
                (
                    "insert",
                    "For this PP-Configuration, a distributed TOE consists of multiple separately deployed TOE components that collectively provide the TOE security functionality. A TOE component is a separately deployed portion of the TOE that is identified in the ST and mapped to the base cPP and relevant SFRs. Each TOE component shall be identified in the ST and mapped to the base cPP and, where applicable, to the Server Module, Agent Module, or both according to the role or roles performed by that TOE component.",
                )
            ],
            [
                (
                    "insert",
                    "For containerized or microservices TOEs, the TOE consists of the application payload components identified in the ST. The container orchestration platform, container runtime, operating system, service mesh infrastructure, ingress infrastructure, cluster networking, and platform-provided secret or configuration stores are part of the operational environment unless explicitly included in the TOE boundary.",
                )
            ],
            [
                (
                    "insert",
                    "The ST shall provide an SFR allocation rationale that identifies whether each claimed requirement is satisfied by all TOE components, by applicable TOE components that perform the relevant function, by at least one TOE component, as a feature dependent requirement, or by an allowed operational environment dependency. The ST shall describe all inter-component TOE communications and identify the mechanisms used to authorize and protect those communications.",
                )
            ],
        ],
    ),
]

config_changes.append(
    (
        "SFR Allocation Tables",
        source_block_as_insertions(
            "Archive/Modules/Agent/appSW_PP_Config_ServerAgent.adoc",
            "==== SFR Allocation for Distributed TOEs",
            "\n== Conformance Claims",
        ),
    )
)


def main():
    build_doc(
        "application-v2-distributed-microservices-redline.docx",
        "Application Software PP v2 Distributed/Microservices Redline",
        "input/application.xml",
        base_changes,
    )
    build_doc(
        "server-module-distributed-microservices-redline.docx",
        "Server Module Distributed/Microservices Redline",
        "Modules/Server/cPP_MOD-Server.adoc",
        server_changes,
    )
    build_doc(
        "agent-module-distributed-microservices-redline.docx",
        "Agent Module Distributed/Microservices Redline",
        "Modules/Agent/cPP_MOD-Agent.adoc",
        agent_changes,
    )
    build_doc(
        "server-agent-configuration-distributed-microservices-redline.docx",
        "Server + Agent PP-Configuration Distributed/Microservices Redline",
        "Modules/Agent/appSW_PP_Config_ServerAgent.adoc",
        config_changes,
    )


if __name__ == "__main__":
    main()
